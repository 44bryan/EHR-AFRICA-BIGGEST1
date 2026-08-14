"""Generate demo_accounts.xlsx from already-created demo accounts."""
import os, django
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'ehr_project.settings.base')
django.setup()

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from apps.accounts.models import User, UserRole

PASSWORD    = 'Micei@2021@@'
EMAIL_DOMAIN = 'demo.gludd.cm'

SLUGS = [
    'super.admin','e.mbarga','c.fouda','p.atanga','s.nkeng',
    'a.betrand','j.noa','m.essomba','p.tchoua','g.nnomo',
    's.mvondo','c.abena','e.biyong','r.owona','f.manga',
    'd.mba','k.simo','l.ndjock','b.tsimi','h.nkoa',
    's.etoa','j.bebe','f.ngono','r.amvam','a.minlo','m.onana',
    'd.mengue','g.eboue','v.mballa','t.nkolo','h.mekongo',
    'c.ndongo','b.zoua','r.etoundi','o.ndoumbe','l.mfou',
    'n.ekwalla','j.nkemba','b.amougou','s.nlend','l.binyom',
    's.mekuate','it.admin',
]

NAMES = {
    'super.admin': 'Super Admin',
    'e.mbarga':    'Emmanuel Mbarga',
    'c.fouda':     'Christelle Fouda',
    'p.atanga':    'Patrick Atanga',
    's.nkeng':     'Dr. Samuel Nkeng',
    'a.betrand':   'Dr. Alice Betrand',
    'j.noa':       'Dr. Jean Noa',
    'm.essomba':   'Dr. Marie Essomba',
    'p.tchoua':    'Dr. Paul Tchoua',
    'g.nnomo':     'Dr. Grace Nnomo',
    's.mvondo':    'Dr. Simon Mvondo',
    'c.abena':     'Dr. Claire Abena',
    'e.biyong':    'Dr. Eric Biyong',
    'r.owona':     'Dr. Ruth Owona',
    'f.manga':     'Dr. Felix Manga',
    'd.mba':       'Dr. Diane Mba',
    'k.simo':      'Dr. Kevin Simo',
    'l.ndjock':    'Dr. Laure Ndjock',
    'b.tsimi':     'Bernard Tsimi',
    'h.nkoa':      'Helene Nkoa',
    's.etoa':      'Sylvie Etoa',
    'j.bebe':      'Josephine Bebe',
    'f.ngono':     'Florence Ngono',
    'r.amvam':     'Rodrigue Amvam',
    'a.minlo':     'Angele Minlo',
    'm.onana':     'Martine Onana',
    'd.mengue':    'Denis Mengue',
    'g.eboue':     'Gaston Eboue',
    'v.mballa':    'Valerie Mballa',
    't.nkolo':     'Thierry Nkolo',
    'h.mekongo':   'Herve Mekongo',
    'c.ndongo':    'Cecile Ndongo',
    'b.zoua':      'Blaise Zoua',
    'r.etoundi':   'Raissa Etoundi',
    'o.ndoumbe':   'Ornella Ndoumbe',
    'l.mfou':      'Leopold Mfou',
    'n.ekwalla':   'Nadege Ekwalla',
    'j.nkemba':    'Josue Nkemba',
    'b.amougou':   'Bertrand Amougou',
    's.nlend':     'Serge Nlend',
    'l.binyom':    'Lydie Binyom',
    's.mekuate':   'Sandra Mekuate',
    'it.admin':    'IT Administrator',
}

ROLE_CHOICES = {v: str(label) for v, label in UserRole.choices}

SECTION_COLORS = {
    'Executive / Leadership':    'EFF6FF',
    'Medical Direction':         'F0FDF4',
    'Doctors / Specialists':     'F8FAFC',
    'Allied Clinical':           'FFF7ED',
    'Nursing':                   'FFF1F2',
    'Technical / Paramedical':   'F5F3FF',
    'Administrative / Finance':  'FFFBEB',
}

ROLE_TO_SECTION = {
    'super_admin': 'Executive / Leadership',
    'ceo': 'Executive / Leadership',
    'finance_director': 'Executive / Leadership',
    'admin_director': 'Executive / Leadership',
    'medical_director': 'Medical Direction',
    'ophthalmologist': 'Doctors / Specialists',
    'retina_specialist': 'Doctors / Specialists',
    'glaucoma_specialist': 'Doctors / Specialists',
    'cataract_surgeon': 'Doctors / Specialists',
    'cornea_specialist': 'Doctors / Specialists',
    'pediatric_ophthalmologist': 'Doctors / Specialists',
    'oculoplastics_surgeon': 'Doctors / Specialists',
    'neuro_ophthalmologist': 'Doctors / Specialists',
    'uveitis_specialist': 'Doctors / Specialists',
    'optometrist': 'Doctors / Specialists',
    'low_vision_specialist': 'Doctors / Specialists',
    'intern': 'Doctors / Specialists',
    'pediatric_refractionist': 'Doctors / Specialists',
    'anaesthetist': 'Allied Clinical',
    'orthoptist': 'Allied Clinical',
    'nurse': 'Nursing',
    'consultation_nurse': 'Nursing',
    'ward_nurse': 'Nursing',
    'theatre_nurse': 'Nursing',
    'counseling_nurse': 'Nursing',
    'nurse_director': 'Nursing',
    'investigation_technician': 'Technical / Paramedical',
    'lab_technician': 'Technical / Paramedical',
    'pharmacist': 'Technical / Paramedical',
    'optical_dispenser': 'Technical / Paramedical',
    'cssd_technician': 'Technical / Paramedical',
    'receptionist': 'Administrative / Finance',
    'admission_officer': 'Administrative / Finance',
    'cashier': 'Administrative / Finance',
    'chief_cashier': 'Administrative / Finance',
    'accountant': 'Administrative / Finance',
    'senior_accountant': 'Administrative / Finance',
    'billing_officer': 'Administrative / Finance',
    'gatekeeper': 'Administrative / Finance',
    'medical_records': 'Administrative / Finance',
    'quality_officer': 'Administrative / Finance',
    'admin': 'Administrative / Finance',
}

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Demo Accounts'

# ── Banner row 1
ws.merge_cells('A1:H1')
cell = ws['A1']
cell.value = 'GLUDD EYE HOSPITAL  --  DEMO ACCOUNTS FOR PRESENTATION'
cell.font = Font(name='Calibri', bold=True, size=14, color='FFFFFF')
cell.fill = PatternFill('solid', fgColor='1D4ED8')
cell.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 34

# ── Password banner row 2
ws.merge_cells('A2:H2')
cell2 = ws['A2']
cell2.value = 'ALL ACCOUNTS USE THE SAME PASSWORD:   ' + PASSWORD + '   --   For authorised demo use only'
cell2.font = Font(name='Calibri', bold=True, size=11, color='1E3A5F')
cell2.fill = PatternFill('solid', fgColor='DBEAFE')
cell2.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[2].height = 22

# ── Column headers row 3
headers = ['#', 'Full Name', 'Login Email', 'Username', 'Password', 'Role / Position', 'Dept Code', 'Emp. ID']
widths  = [5,   26,          34,            22,          20,         34,                 14,           12]
hfill = PatternFill('solid', fgColor='1E3A5F')
hfont = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
for ci, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(row=3, column=ci, value=h)
    c.font = hfont
    c.fill = hfill
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.column_dimensions[get_column_letter(ci)].width = w
ws.row_dimensions[3].height = 22

tb = Border(
    left=Side(style='thin', color='E5E7EB'),
    right=Side(style='thin', color='E5E7EB'),
    bottom=Side(style='thin', color='E5E7EB'),
)

# ── Data rows
row_num = 4
for idx, slug in enumerate(SLUGS, 1):
    email = slug + '@' + EMAIL_DOMAIN
    try:
        u = User.objects.get(email=email)
    except User.DoesNotExist:
        print('MISSING:', email)
        continue

    role_label = ROLE_CHOICES.get(u.role, u.role)
    section    = ROLE_TO_SECTION.get(u.role, 'Administrative / Finance')
    bg         = SECTION_COLORS.get(section, 'FFFFFF')
    rfill      = PatternFill('solid', fgColor=bg)
    dept_code  = u.department.code if u.department else '--'

    values = [
        str(idx),
        NAMES.get(slug, u.get_full_name()),
        email,
        slug,
        PASSWORD,
        role_label,
        dept_code,
        str(u.employee_id),
    ]

    for ci, val in enumerate(values, 1):
        c = ws.cell(row=row_num, column=ci, value=val)
        c.fill   = rfill
        c.border = tb
        c.alignment = Alignment(vertical='center')
        if ci == 1:
            c.font = Font(name='Calibri', bold=True, size=9)
            c.alignment = Alignment(horizontal='center', vertical='center')
        elif ci == 5:
            c.font = Font(name='Courier New', bold=True, size=10, color='DC2626')
        elif ci == 3:
            c.font = Font(name='Calibri', size=9, color='1D4ED8')
        else:
            c.font = Font(name='Calibri', size=9)

    ws.row_dimensions[row_num].height = 17
    row_num += 1

# ── Footer
ws.merge_cells('A' + str(row_num) + ':H' + str(row_num))
fc = ws.cell(row=row_num, column=1,
             value='CONFIDENTIAL  --  ' + str(len(SLUGS)) + ' demo accounts  --  Password: ' + PASSWORD)
fc.font = Font(name='Calibri', italic=True, size=9, color='6B7280')
fc.fill = PatternFill('solid', fgColor='F9FAFB')
fc.alignment = Alignment(horizontal='center')

ws.freeze_panes = 'A4'

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'demo_accounts.xlsx')
wb.save(out)
print('Excel saved to:', out)
